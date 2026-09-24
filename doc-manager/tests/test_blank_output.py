"""A document with no readable text must never be reported as a success.

An image-only scan makes the OCR pipeline return empty page text; build_docx() still
produces a valid ~36KB Word file, so the task used to be marked "done" with the user's
quota spent and not one visible character inside. Three such pages gave a file whose
entire content was "-- Page 2 ---- Page 3 --".
"""

import io

import pytest
from docx import Document
from prometheus_client import REGISTRY

from app.core.config import settings
from app.worker import consumer

ARABIC = "هذا نص عربي حقيقي من الصفحة الممسوحة"


def sample(name, **labels):
    return REGISTRY.get_sample_value(name, labels) or 0.0


def make_docx(pages: list[str]) -> bytes:
    """A DOCX shaped like the pipeline's output, built here rather than imported.

    Deliberately not calling app.ocr.pipeline.build_docx: the guard reads the produced
    file, not the pipeline, and these tests should keep passing while the pipeline is
    rewritten. test_real_pipeline_blank_output_is_detected covers the real builder.
    """
    doc = Document()
    for n, text in enumerate(pages, 1):
        if n > 1:
            doc.add_paragraph(f"-- Page {n} --")
        for line in text.split("\n"):
            doc.add_paragraph(line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def run(storage, monkeypatch):
    """Run one task whose OCR backend returns the given DOCX bytes."""
    refunds: list[tuple] = []
    monkeypatch.setattr(consumer, "OCR_BACKEND", "gemini")
    monkeypatch.setattr(
        consumer, "_refund_pages", lambda *a, **k: refunds.append((a, k))
    )

    def _run(docx_bytes, *, user_id="42", file_id="a.pdf", pages=3):
        monkeypatch.setattr(consumer, "_call_gemini", lambda *a, **k: docx_bytes)
        storage.save_file(user_id, file_id, io.BytesIO(b"%PDF-1.4 fake"))
        consumer.process_task(
            {"file_id": file_id, "user_id": user_id, "pages": pages, "mode": "ocr"}
        )
        return storage.get_status(user_id, file_id)

    _run.refunds = refunds
    return _run


# ── the reported failure ──────────────────────────────────────────────────────

def test_document_with_no_text_is_failed_not_done(run, storage):
    empty0 = sample("ocr_requests_total", status="empty", mode="ocr")
    ok0 = sample("ocr_requests_total", status="success", mode="ocr")
    err0 = sample("ocr_requests_total", status="error", mode="ocr")

    assert run(make_docx(["", "", ""])) == "failed"

    assert run.refunds, "the reserved pages must be refunded"
    assert run.refunds[0][0][:3] == ("42", 3, None)
    # Counted apart from a crash: "we ship blank documents" and "the pipeline is
    # throwing" are different incidents.
    assert sample("ocr_requests_total", status="empty", mode="ocr") == empty0 + 1
    assert sample("ocr_requests_total", status="success", mode="ocr") == ok0
    assert sample("ocr_requests_total", status="error", mode="ocr") == err0


def test_blank_output_is_never_saved(run, storage):
    run(make_docx(["", ""]))
    assert not storage.file_exists("42", "a.docx")


def test_page_separators_alone_do_not_count_as_text(run):
    """The exact shape of the bug report: the only text is "-- Page 2 ---- Page 3 --"."""
    assert consumer._docx_visible_chars(make_docx(["", "", ""])) == 0


def test_real_pipeline_blank_output_is_detected():
    """Against the actual builder, so the guard cannot drift from what ships.

    Skipped rather than failed if build_docx's signature moves — the pipeline is
    owned elsewhere, and this file's other tests already pin the behaviour.
    """
    try:
        from app.ocr.pipeline import build_docx

        buf = io.BytesIO()
        build_docx(["", "", ""], buf)
        blank = buf.getvalue()
    except Exception as e:  # noqa: BLE001
        pytest.skip(f"pipeline.build_docx not callable as (pages, out): {e}")

    assert len(blank) > 1000, "a blank DOCX is still a large, valid file"
    assert consumer._docx_visible_chars(blank) < settings.OCR_MIN_OUTPUT_CHARS


# ── the negative cases: do not fail good work ─────────────────────────────────

def test_good_document_is_still_done(run, storage):
    ok0 = sample("ocr_requests_total", status="success", mode="ocr")
    assert run(make_docx([ARABIC, ARABIC])) == "done"
    assert run.refunds == []
    assert storage.file_exists("42", "a.docx")
    assert sample("ocr_requests_total", status="success", mode="ocr") == ok0 + 1


def test_blank_page_in_the_middle_does_not_fail_the_document(run):
    """A title page, a blank verso, a full-page figure: normal in a real book.

    The rule is document-wide on purpose — a per-page rule would throw away 39 good
    pages because page 12 was a photograph.
    """
    pages = [""] + [ARABIC] * 38 + ["", ARABIC, ""]
    assert run(make_docx(pages)) == "done"
    assert run.refunds == []


def test_one_page_of_text_carries_a_mostly_blank_document(run):
    assert run(make_docx(["", "", ARABIC, ""])) == "done"


def test_single_artifact_character_is_still_blank(run):
    """OCR that returns "." or a stray digit per page is not a conversion."""
    assert run(make_docx([".", "-", "1"])) == "failed"


# ── the guard itself ──────────────────────────────────────────────────────────

def test_unreadable_output_is_not_assumed_blank(run):
    """If the output stops being a DOCX we must not fail every conversion.

    A wrong "empty" verdict here would take the whole product down; the correct
    reading of unparseable bytes is "some other bug", not "no text".
    """
    assert consumer._docx_visible_chars(b"PK\x03\x04 not really a docx") is None
    assert run(b"PK\x03\x04 not really a docx") == "done"


def test_threshold_is_configurable(run, monkeypatch):
    monkeypatch.setattr(settings, "OCR_MIN_OUTPUT_CHARS", 500)
    assert run(make_docx([ARABIC])) == "failed"
    monkeypatch.setattr(settings, "OCR_MIN_OUTPUT_CHARS", 0)
    assert run(make_docx([""])) == "done"


def test_visible_chars_ignores_whitespace_and_markers(run):
    assert consumer._docx_visible_chars(make_docx(["   \n  \n   "])) == 0
    assert consumer._docx_visible_chars(make_docx(["ab cd"])) == 4
    # A separator written some other way is still chrome, not content.
    assert consumer._docx_visible_chars(make_docx(["— Page 2 —", "## page 3"])) == 0
