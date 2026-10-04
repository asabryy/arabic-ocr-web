"""A DOCX builder that carries the source document's styling.

``pipeline.build_docx`` normalises everything to one font and one size, because
its input is a vision model's plain text and carries no styling to preserve. The
digital path reads real font names, point sizes, colours and weights off the
glyphs, so it needs a builder that can write them.

Shares every OOXML primitive with the Gemini path (``app.ocr.ooxml``) and the
same direction logic (``app.ocr.script``); only the run properties differ.
"""

import io
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_DIRECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.ocr.ooxml import (
    A4_HEIGHT_TWIPS,
    A4_WIDTH_TWIPS,
    PPR_SEQ,
    RPR_SEQ,
    SECTPR_SEQ,
    add_break,
    half_points,
    harden_rtl,
    set_child,
    twips,
    xml_safe,
)
from app.ocr.script import is_rtl_text, segment_by_script

STYLE_SOURCE = "source"
STYLE_UNIFORM = "uniform"


# ── Model ────────────────────────────────────────────────────────────────────


@dataclass
class Run:
    """One stretch of text sharing a single appearance."""

    text: str
    size: float = 12.0
    color: int | None = None       # 0xRRGGBB, as PyMuPDF reports it
    bold: bool = False
    italic: bool = False
    underline: bool = False
    font: str = "Arial"


@dataclass
class PageBorder:
    style: str = "single"          # single | double | thickThinSmallGap | ...
    size_eighths: int = 8          # w:sz, in eighths of a point (2-96)
    space_pt: int = 24             # w:space, offset from the page edge (0-31)
    color: str = "000000"


@dataclass
class PageSetup:
    """Everything that belongs to the section rather than a paragraph."""

    margin_top_pt: float = 72.0
    margin_bottom_pt: float = 72.0
    margin_left_pt: float = 72.0
    margin_right_pt: float = 72.0
    border: PageBorder | None = None
    page_number_align: str | None = None    # start | center | end
    page_number_start: str | None = None


@dataclass
class Block:
    """One renderable unit of a page."""

    kind: str                      # para | heading | bullet | ordered | table
    runs: list = field(default_factory=list)
    level: int = 0
    align: str | None = None       # start | end | center | both
    rows: list = field(default_factory=list)   # list[list[list[Run]]]
    has_header_row: bool = False


# ── Font mapping ─────────────────────────────────────────────────────────────

# Subset prefixes look like "ABCDEF+Arial-BoldMT" and mean nothing to Word.
_STYLE_SUFFIXES = (
    "-boldmt", "-bolditalicmt", "-italicmt", "-bold", "-italic", "-regular",
    "-boldoblique", "-oblique", "psmt", "mt", "ps",
)


def map_font(pdf_font: str, fallback: str = "Arial") -> str:
    """Turn an embedded PDF font name into one Word can resolve.

    Weight and slope live in ``w:b``/``w:i``, not in the family name, so the
    suffixes encoding them are stripped: "ABCDEF+Arial-BoldMT" -> "Arial".
    """
    if not pdf_font:
        return fallback
    name = pdf_font.split("+", 1)[-1]
    low = name.lower()
    for suffix in _STYLE_SUFFIXES:
        if low.endswith(suffix):
            name = name[: len(name) - len(suffix)]
            low = name.lower()
            break
    name = name.replace("-", " ").strip()
    return name or fallback


# ── Runs ─────────────────────────────────────────────────────────────────────


def _add_run(para, text: str, run: Run, *, rtl: bool):
    """One Word run carrying BOTH the Latin and complex-script properties.

    Word lays Arabic out from ``w:cs``/``w:szCs``, so a property set only on the
    Latin side silently does nothing — the same reason ``pipeline._add_run``
    writes the XML directly. Extended here with colour, underline and the real
    family name.
    """
    r = para.add_run(xml_safe(text))
    rPr = r._r.get_or_add_rPr()

    font = run.font or "Arial"
    set_child(rPr, RPR_SEQ, "w:rFonts",
              **{"w:ascii": font, "w:hAnsi": font, "w:cs": font})

    if run.bold:
        # w:bCs is the complex-script half of bold; w:b alone leaves Arabic plain.
        set_child(rPr, RPR_SEQ, "w:b")
        set_child(rPr, RPR_SEQ, "w:bCs")
    if run.italic:
        set_child(rPr, RPR_SEQ, "w:i")
        set_child(rPr, RPR_SEQ, "w:iCs")

    set_child(rPr, RPR_SEQ, "w:sz", **{"w:val": half_points(run.size)})
    set_child(rPr, RPR_SEQ, "w:szCs", **{"w:val": half_points(run.size)})

    if run.color is not None:
        set_child(rPr, RPR_SEQ, "w:color", **{"w:val": f"{run.color:06X}"})
    if run.underline:
        set_child(rPr, RPR_SEQ, "w:u", **{"w:val": "single"})
    if rtl:
        set_child(rPr, RPR_SEQ, "w:rtl")
    return r


def _emit_runs(para, runs: list[Run], default_rtl: bool) -> None:
    """Write runs, splitting each at script boundaries so Latin and digits inside
    an Arabic line keep their own direction."""
    for run in runs:
        if not run.text:
            continue
        for seg, seg_rtl in segment_by_script(run.text, default_rtl):
            if seg:
                _add_run(para, seg, run, rtl=seg_rtl)


def _direct_paragraph(para, rtl: bool, align: str | None) -> None:
    """Base direction plus LOGICAL alignment.

    ``start``/``end`` rather than ``left``/``right``: Word for Mac resolves
    ``right`` as the logical end of the line, which in an RTL paragraph is the
    visually left margin.
    """
    pPr = para._p.get_or_add_pPr()
    set_child(pPr, PPR_SEQ, "w:bidi", **{"w:val": "1" if rtl else "0"})
    set_child(pPr, PPR_SEQ, "w:jc", **{"w:val": align or ("end" if rtl else "start")})


# ── Blocks ───────────────────────────────────────────────────────────────────

_LIST_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")


def _text_of(runs: list[Run]) -> str:
    return "".join(r.text for r in runs)


def _render_table(doc, block: Block, default_rtl: bool) -> None:
    rows = block.rows
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    try:
        table.style = "Table Grid"
    except KeyError:  # pragma: no cover — present in the default template
        pass

    flat = " ".join(_text_of(cell) for row in rows for cell in row)
    table_rtl = is_rtl_text(flat, default=default_rtl)
    if table_rtl:
        # Without bidiVisual Word draws column 1 on the LEFT, mirroring the source.
        table.table_direction = WD_TABLE_DIRECTION.RTL

    for r_i, row in enumerate(rows):
        for c_i in range(width):
            cell_runs = row[c_i] if c_i < len(row) else []
            para = table.cell(r_i, c_i).paragraphs[0]
            cell_rtl = is_rtl_text(_text_of(cell_runs), default=table_rtl)
            _direct_paragraph(para, cell_rtl, None)
            _emit_runs(para, cell_runs, cell_rtl)
    doc.add_paragraph()


def _render_block(doc, block: Block, default_rtl: bool) -> None:
    if block.kind == "table":
        _render_table(doc, block, default_rtl)
        return

    text = _text_of(block.runs)
    if not text.strip():
        return
    rtl = is_rtl_text(text, default=default_rtl)

    if block.kind == "heading":
        para = doc.add_paragraph(style=f"Heading {min(max(block.level, 1), 2)}")
    elif block.kind == "bullet":
        para = doc.add_paragraph(style=_LIST_BULLET_STYLES[min(block.level, 2)])
    elif block.kind == "ordered":
        para = doc.add_paragraph(style="List Paragraph")
    else:
        para = doc.add_paragraph()

    _direct_paragraph(para, rtl, block.align)
    _emit_runs(para, block.runs, rtl)


# ── Section ──────────────────────────────────────────────────────────────────


def _apply_border(sectPr, border: PageBorder) -> None:
    """A frame on all four sides.

    ``offsetFrom="page"`` measures the inset from the paper edge, which is how
    producers draw it; the default ("text") measures from the margin and puts the
    frame in the wrong place on any document with wide margins.
    """
    el = set_child(sectPr, SECTPR_SEQ, "w:pgBorders", **{"w:offsetFrom": "page"})
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), border.style)
        e.set(qn("w:sz"), str(max(2, min(96, border.size_eighths))))
        e.set(qn("w:space"), str(max(0, min(31, int(round(border.space_pt))))))
        e.set(qn("w:color"), border.color)
        el.append(e)


def _add_page_field(container, align: str, rtl: bool) -> None:
    """A real PAGE field, so Word renumbers as the document is edited."""
    paragraphs = container.paragraphs
    para = paragraphs[0] if paragraphs else container.add_paragraph()
    para.clear()
    _direct_paragraph(para, rtl, align)
    placeholder = _add_run(para, "1", Run("1", size=9.0), rtl=False)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    fld.append(placeholder._r)      # re-parents the run into the field
    para._p.append(fld)


def _configure_section(section, default_rtl: bool, setup: PageSetup) -> None:
    sectPr = section._sectPr
    set_child(sectPr, SECTPR_SEQ, "w:pgSz",
              **{"w:w": str(A4_WIDTH_TWIPS), "w:h": str(A4_HEIGHT_TWIPS)})
    set_child(sectPr, SECTPR_SEQ, "w:pgMar", **{
        "w:top": twips(setup.margin_top_pt),
        "w:right": twips(setup.margin_right_pt),
        "w:bottom": twips(setup.margin_bottom_pt),
        "w:left": twips(setup.margin_left_pt),
        "w:header": twips(min(setup.margin_top_pt / 2, 36)),
        "w:footer": twips(min(setup.margin_bottom_pt / 2, 36)),
        "w:gutter": "0",
    })
    if setup.border:
        _apply_border(sectPr, setup.border)
    if default_rtl:
        set_child(sectPr, SECTPR_SEQ, "w:bidi", **{"w:val": "1"})


# ── Entry point ──────────────────────────────────────────────────────────────


def build_styled_docx(pages: list[list[Block]], out, setup: PageSetup | None = None) -> None:
    """Pages of styled blocks -> a Word document that keeps the source's styling."""
    setup = setup or PageSetup()

    all_text = " ".join(
        _text_of(b.runs) + " " + " ".join(_text_of(c) for row in b.rows for c in row)
        for page in pages for b in page
    )
    default_rtl = is_rtl_text(all_text, default=True)

    doc = Document()
    section = doc.sections[0]
    _configure_section(section, default_rtl, setup)

    start = setup.page_number_start
    if start and start.isdigit() and 0 < int(start) < 32768:
        set_child(section._sectPr, SECTPR_SEQ, "w:pgNumType", **{"w:start": start})
    if setup.page_number_align:
        _add_page_field(section.footer, setup.page_number_align, default_rtl)

    for index, blocks in enumerate(pages):
        if index:
            add_break(doc.add_paragraph(), "page")
        for block in blocks:
            _render_block(doc, block, default_rtl)

    # Same hardening as the Gemini path: the two RTL layers python-docx cannot
    # express are applied to the saved package.
    buf = io.BytesIO()
    doc.save(buf)
    hardened = harden_rtl(buf.getvalue())
    if hasattr(out, "write"):
        out.write(hardened)
    else:
        with open(out, "wb") as fh:
            fh.write(hardened)
